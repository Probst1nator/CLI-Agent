"""
ReadFile Tool Tests

Tests for the readfile tool including basic reading,
truncation handling, line-indexed access, and error handling.
"""

import os
import sys
import tempfile
from pathlib import Path

# Setup test imports
test_root = Path(__file__).parent.parent.parent
sys.path.insert(0, str(test_root))

# Import the tool
from tools.readfile import readfile


def test_readfile_basic_reading():
    """Test basic file reading functionality"""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "test.txt")

        # Create test file with known content
        content = """Line 1: Hello World
Line 2: This is a test
Line 3: With multiple lines
Line 4: End of file"""

        with open(test_file, 'w', encoding='utf-8') as f:
            f.write(content)

        # Test basic reading
        xml_input = f"<filepath>{test_file}</filepath>"
        result = readfile.run(xml_input)

        # Should return the exact content
        assert result == content, f"Content mismatch. Expected: {repr(content)}, Got: {repr(result)}"

        return True


def test_readfile_line_indexed_reading():
    """Test line-indexed reading functionality"""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "indexed_test.txt")

        # Create file with numbered lines
        lines = [f"Line {i+1}: Content for line {i+1}" for i in range(20)]
        content = "\n".join(lines)

        with open(test_file, 'w', encoding='utf-8') as f:
            f.write(content)

        # Test single line reading
        xml_input = f"<filepath>{test_file}</filepath><lines>5</lines>"
        result = readfile.run(xml_input)
        expected = "Line 5: Content for line 5"
        assert result == expected, f"Single line read failed. Expected: {repr(expected)}, Got: {repr(result)}"

        # Test line range reading
        xml_input = f"<filepath>{test_file}</filepath><lines>3-6</lines>"
        result = readfile.run(xml_input)
        expected = "\n".join(lines[2:6])  # Lines 3-6 (0-indexed 2-5)
        assert result == expected, f"Line range read failed. Expected: {repr(expected)}, Got: {repr(result)}"

        # Test reading from line to end
        xml_input = f"<filepath>{test_file}</filepath><lines>18-</lines>"
        result = readfile.run(xml_input)
        expected = "\n".join(lines[17:])  # Lines 18 to end
        assert result == expected, f"Read to end failed. Expected: {repr(expected)}, Got: {repr(result)}"

        # Test reading from start to line
        xml_input = f"<filepath>{test_file}</filepath><lines>-5</lines>"
        result = readfile.run(xml_input)
        expected = "\n".join(lines[:5])  # Lines 1-5 (0-indexed 0-4)
        assert result == expected, f"Read from start failed. Expected: {repr(expected)}, Got: {repr(result)}"

        return True


def test_readfile_large_file_truncation():
    """Test truncation behavior for large files"""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "large_test.txt")

        # Create a file larger than CHUNK_LIMIT
        line = "This is a line with some content to make the file large.\n"
        # Make file larger than the 50000 character CHUNK_LIMIT
        large_content = line * 2000  # Should be well over 50k characters

        with open(test_file, 'w', encoding='utf-8') as f:
            f.write(large_content)

        # First read should be truncated
        xml_input = f"<filepath>{test_file}</filepath>"
        result = readfile.run(xml_input)

        # Should contain truncation message
        assert "File is too large" in result, "Should indicate truncation"
        assert "Truncating to show start and end" in result, "Should explain truncation"
        assert len(result) < len(large_content), "Result should be shorter than original"

        # Should contain content from start and end
        assert large_content[:100] in result, "Should contain content from start"
        assert large_content[-100:] in result, "Should contain content from end"

        return True


def test_readfile_error_cases():
    """Test various error conditions"""
    with tempfile.TemporaryDirectory() as temp_dir:

        # Test non-existent file
        fake_file = os.path.join(temp_dir, "nonexistent.txt")
        xml_input = f"<filepath>{fake_file}</filepath>"
        result = readfile.run(xml_input)
        assert "Error: File not found" in result, "Should fail for non-existent file"

        # Test invalid line specification
        test_file = os.path.join(temp_dir, "test.txt")
        with open(test_file, 'w', encoding='utf-8') as f:
            f.write("line 1\nline 2\nline 3")

        xml_input = f"<filepath>{test_file}</filepath><lines>invalid</lines>"
        result = readfile.run(xml_input)
        assert "Error: Invalid line specification" in result, "Should fail with invalid line spec"

        # Test line number exceeding file length
        xml_input = f"<filepath>{test_file}</filepath><lines>10</lines>"
        result = readfile.run(xml_input)
        assert "Error extracting lines" in result or "Error: Invalid line specification" in result, "Should fail for line beyond file length"

        # Test invalid line range
        xml_input = f"<filepath>{test_file}</filepath><lines>5-2</lines>"
        result = readfile.run(xml_input)
        assert "Error extracting lines" in result or "Error: Invalid line specification" in result, "Should fail for invalid range"

        return True


def test_readfile_empty_file():
    """Test reading empty files"""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "empty.txt")

        # Create empty file
        with open(test_file, 'w', encoding='utf-8') as f:
            f.write("")

        xml_input = f"<filepath>{test_file}</filepath>"
        result = readfile.run(xml_input)

        assert result == "", f"Empty file should return empty string, got: {repr(result)}"

        return True


def test_readfile_special_characters():
    """Test reading files with special characters"""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "special.txt")

        # Content with various encodings and special characters
        content = """Unicode test: ñáéíóú ☀️ 🚀
Quotes: "double" 'single' "smart" 'smart'
Symbols: @#$%^&*()_+-=[]{}|;:,.<>?
Tabs:\tand\tnewlines
"""

        with open(test_file, 'w', encoding='utf-8') as f:
            f.write(content)

        xml_input = f"<filepath>{test_file}</filepath>"
        result = readfile.run(xml_input)

        assert result == content, f"Special characters not preserved. Expected: {repr(content)}, Got: {repr(result)}"

        return True


def test_readfile_different_line_endings():
    """Test reading files with different line endings"""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "line_endings.txt")

        # Create file with Unix line endings
        content = "Line 1\nLine 2\nLine 3"

        with open(test_file, 'w', encoding='utf-8') as f:
            f.write(content)

        xml_input = f"<filepath>{test_file}</filepath>"
        result = readfile.run(xml_input)

        assert result == content, f"Line endings not preserved. Expected: {repr(content)}, Got: {repr(result)}"

        return True


def test_readfile_relative_paths():
    """Test reading with relative paths"""
    original_cwd = os.getcwd()

    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            os.chdir(temp_dir)

            # Create file in current directory
            test_content = "Relative path test content"
            with open("relative_test.txt", 'w', encoding='utf-8') as f:
                f.write(test_content)

            # Test relative path reading
            xml_input = "<filepath>./relative_test.txt</filepath>"
            result = readfile.run(xml_input)

            assert result == test_content, f"Relative path reading failed. Expected: {repr(test_content)}, Got: {repr(result)}"

    finally:
        os.chdir(original_cwd)

    return True


def test_readfile_whitespace_handling():
    """Test reading files with various whitespace patterns"""
    with tempfile.TemporaryDirectory() as temp_dir:
        test_file = os.path.join(temp_dir, "whitespace.txt")

        # Content with leading/trailing whitespace and empty lines
        content = """  Line with leading spaces
\tLine with leading tab
Line with trailing spaces
\t\tMultiple tabs and spaces  \t

\t
Line after empty lines"""

        with open(test_file, 'w', encoding='utf-8') as f:
            f.write(content)

        xml_input = f"<filepath>{test_file}</filepath>"
        result = readfile.run(xml_input)

        assert result == content, f"Whitespace not preserved. Expected: {repr(content)}, Got: {repr(result)}"

        return True


def test_readfile_office_documents():
    """Test reading office document formats (.docx, .xlsx, .pptx)"""
    try:
        import docx
        import openpyxl
        import pptx
    except ImportError:
        # Skip test if libraries not installed
        print("Skipping office document tests - required libraries not installed")
        return True

    with tempfile.TemporaryDirectory() as temp_dir:
        # Test DOCX
        docx_file = os.path.join(temp_dir, "test.docx")
        doc = docx.Document()
        doc.add_heading('Test Document', 0)
        doc.add_paragraph('This is test content.')
        doc.save(docx_file)

        result = readfile.run(docx_file)
        assert "Test Document" in result, "DOCX: Should extract heading"
        assert "This is test content" in result, "DOCX: Should extract paragraph"
        assert "Word document" in result, "DOCX: Should identify as Word document"

        # Test XLSX
        xlsx_file = os.path.join(temp_dir, "test.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws['A1'] = 'Header1'
        ws['B1'] = 'Header2'
        ws['A2'] = 'Data1'
        ws['B2'] = 'Data2'
        wb.save(xlsx_file)

        result = readfile.run(xlsx_file)
        assert "Header1" in result, "XLSX: Should extract cell content"
        assert "Data1" in result, "XLSX: Should extract data"
        assert "Excel spreadsheet" in result, "XLSX: Should identify as Excel"

        # Test PPTX
        pptx_file = os.path.join(temp_dir, "test.pptx")
        prs = pptx.Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = "Test Slide Title"
        prs.save(pptx_file)

        result = readfile.run(pptx_file)
        assert "Test Slide Title" in result, "PPTX: Should extract slide title"
        assert "PowerPoint presentation" in result, "PPTX: Should identify as PowerPoint"

        return True


def test_readfile_office_document_detection():
    """Test office document detection without file extensions"""
    try:
        import docx
    except ImportError:
        print("Skipping office document detection test - python-docx not installed")
        return True

    with tempfile.TemporaryDirectory() as temp_dir:
        # Create DOCX without extension
        docx_file = os.path.join(temp_dir, "no_extension")
        doc = docx.Document()
        doc.add_paragraph('Test content for detection')
        doc.save(docx_file)

        # Should still detect and process it
        is_office, format_type = readfile.is_office_document(docx_file)
        assert is_office, "Should detect office document without extension"
        assert format_type == 'docx', f"Should detect as docx, got {format_type}"

        # Should be able to read it
        result = readfile.run(docx_file)
        assert "Test content for detection" in result, "Should extract content from file without extension"

        return True


def test_readfile_csv_files():
    """Test reading and formatting CSV files."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Create CSV file
        csv_file = os.path.join(temp_dir, "test.csv")
        with open(csv_file, 'w', newline='') as f:
            import csv as csv_module
            writer = csv_module.writer(f)
            writer.writerow(['Name', 'Age', 'City'])
            writer.writerow(['Alice', '25', 'NYC'])
            writer.writerow(['Bob', '30', 'LA'])

        result = readfile.run(csv_file)
        assert "CSV file" in result, "Should identify as CSV"
        assert "Name" in result and "Age" in result, "Should contain headers"
        assert "Alice" in result and "Bob" in result, "Should contain data"
        assert "3 rows" in result, "Should show row count"

        return True


def test_readfile_html_files():
    """Test reading and extracting text from HTML files."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Create HTML file
        html_file = os.path.join(temp_dir, "test.html")
        html_content = """<!DOCTYPE html>
<html>
<head><title>Test</title><script>alert('skip me');</script></head>
<body><h1>Heading</h1><p>Paragraph text</p></body>
</html>"""
        with open(html_file, 'w') as f:
            f.write(html_content)

        result = readfile.run(html_file)
        assert "HTML file" in result, "Should identify as HTML"
        assert "Heading" in result, "Should extract heading"
        assert "Paragraph text" in result, "Should extract paragraph"
        assert "alert" not in result, "Should skip script tags"

        return True


def test_readfile_zip_archives():
    """Test reading ZIP archive contents."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Create ZIP file
        zip_file = os.path.join(temp_dir, "test.zip")
        import zipfile
        with zipfile.ZipFile(zip_file, 'w') as zf:
            zf.writestr('file1.txt', 'Content 1')
            zf.writestr('folder/file2.txt', 'Content 2')

        result = readfile.run(zip_file)
        assert "Archive" in result, "Should identify as archive"
        assert "file1.txt" in result, "Should list first file"
        assert "folder/file2.txt" in result, "Should list nested file"
        assert "2 file(s)" in result, "Should show file count"

        return True


def test_readfile_tar_gz_archives():
    """Test reading TAR.GZ archive contents."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Create TAR.GZ file
        tar_file = os.path.join(temp_dir, "test.tar.gz")
        import io
        import tarfile

        with tarfile.open(tar_file, 'w:gz') as tf:
            file1 = io.BytesIO(b'Tar content 1')
            info1 = tarfile.TarInfo(name='tarfile1.txt')
            info1.size = len(file1.getvalue())
            file1.seek(0)
            tf.addfile(tarinfo=info1, fileobj=file1)

        result = readfile.run(tar_file)
        assert "Archive" in result, "Should identify as archive"
        assert "tarfile1.txt" in result, "Should list tar file"

        return True


def test_readfile_sqlite_databases():
    """Test reading SQLite database schema and data."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Create SQLite database
        db_file = os.path.join(temp_dir, "test.db")
        import sqlite3

        conn = sqlite3.connect(db_file)
        cursor = conn.cursor()
        cursor.execute('CREATE TABLE users (id INTEGER PRIMARY KEY, name TEXT, age INTEGER)')
        cursor.execute('INSERT INTO users VALUES (1, "Alice", 25)')
        cursor.execute('INSERT INTO users VALUES (2, "Bob", 30)')
        conn.commit()
        conn.close()

        result = readfile.run(db_file)
        assert "SQLite database" in result, "Should identify as SQLite"
        assert "users" in result, "Should show table name"
        assert "2 rows" in result, "Should show row count"
        assert "Alice" in result and "Bob" in result, "Should show sample data"
        assert "PRIMARY KEY" in result, "Should show primary key"

        return True


def test_readfile_jupyter_notebooks():
    """Test reading and formatting Jupyter notebooks."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Create Jupyter notebook
        nb_file = os.path.join(temp_dir, "test.ipynb")
        import json

        notebook = {
            "cells": [
                {
                    "cell_type": "markdown",
                    "source": ["# Test Notebook"]
                },
                {
                    "cell_type": "code",
                    "source": ["print('Hello')"],
                    "outputs": [
                        {"output_type": "stream", "text": ["Hello\n"]}
                    ]
                }
            ],
            "metadata": {},
            "nbformat": 4,
            "nbformat_minor": 2
        }

        with open(nb_file, 'w') as f:
            json.dump(notebook, f)

        result = readfile.run(nb_file)
        assert "Jupyter Notebook" in result, "Should identify as Jupyter notebook"
        assert "# Test Notebook" in result, "Should show markdown content"
        assert "print('Hello')" in result, "Should show code content"
        assert "Hello" in result, "Should show output"
        assert "2 cell(s)" in result, "Should show cell count"

        return True


# Test level assignments
test_readfile_basic_reading.level = 2
test_readfile_line_indexed_reading.level = 2
test_readfile_large_file_truncation.level = 3
test_readfile_error_cases.level = 2
test_readfile_empty_file.level = 2
test_readfile_special_characters.level = 2
test_readfile_different_line_endings.level = 2
test_readfile_relative_paths.level = 2
test_readfile_whitespace_handling.level = 2
test_readfile_office_documents.level = 2
test_readfile_office_document_detection.level = 2
test_readfile_csv_files.level = 2
test_readfile_html_files.level = 2
test_readfile_zip_archives.level = 2
test_readfile_tar_gz_archives.level = 2
test_readfile_sqlite_databases.level = 2
test_readfile_jupyter_notebooks.level = 2